"""Wandelt 'Arbeitsstunden_und_Rechnungsstellung_<Jahr>.xlsx' in eine Importdatei fuer die Zeiterfassungs-App um.
Aufruf: python3 excel_to_json.py <excel> <jahr> <ausgabe.json>"""
import openpyxl, sys, json, datetime, re

src, year, out = sys.argv[1], int(sys.argv[2]), sys.argv[3]
wb = openpyxl.load_workbook(src)
wv = openpyxl.load_workbook(src, data_only=True)
ws, wsv = wb['Stunden'], wv['Stunden']

HDR = 17
first_data, last_data = 18, 383
# letzte Projektspalte = Spalte vor 'Effektive h'
last_col = next(c for c in range(2, ws.max_column + 1) if ws.cell(16, c).value == 'Effektive h') - 1

def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None

def date8(v):
    if isinstance(v, (int, float)) and 19000101 < v < 21000101:
        s = str(int(v)); return f"{s[:4]}-{s[4:6]}-{s[6:]}"
    return None

comments = {}
for row in ws.iter_rows(min_row=1, max_row=HDR):
    for c in row:
        if c.comment:
            t = c.comment.text
            t = t.split('Comment:')[-1].strip() if 'Comment:' in t else t.strip()
            comments.setdefault(c.column, []).append(f"{ws.cell(c.row,1).value or 'Projekt'}: {t}")

projects, col2id = [], {}
ff = {1: None, 2: None, 3: None}
for c in range(2, last_col + 1):
    for r in (1, 2, 3):
        v = ws.cell(r, c).value
        if isinstance(v, str) and v.strip():
            ff[r] = v.strip()
            if r == 1: ff[2] = None
    name = (ws.cell(4, c).value or '').strip()
    kunde = ff[3] or ''
    bereich = ff[1] or 'Intern'
    kategorie = ff[2] or ('Intern' if not ff[1] else '')
    mw_raw = (ws.cell(5, c).value or '').strip().lower()
    mwst = 'drauf' if 'schlagen' in mw_raw else 'inkl' if 'inkl' in mw_raw else 'keine'
    betrag, kosten = num(wsv.cell(6, c).value), num(wsv.cell(7, c).value) or 0.0
    notes = list(comments.get(c, []))
    so_raw = wsv.cell(9, c).value
    so = num(so_raw)
    if isinstance(so_raw, str) and so_raw.strip().lower() == 'pauschal':
        notes.append('Stunden (Offerte): pauschal')
    f6 = ws.cell(6, c).value
    if isinstance(f6, str) and f6.startswith('='): notes.append(f"Ansatz (brutto) im Excel: {f6}")
    f7 = ws.cell(7, c).value
    if isinstance(f7, str) and f7.startswith('='): notes.append(f"Kosten im Excel: {f7}")
    rd_raw, ze_raw = wsv.cell(16, c).value, wsv.cell(17, c).value
    rd, ze = date8(rd_raw), date8(ze_raw)
    rds = rd_raw.strip().lower() if isinstance(rd_raw, str) else ''
    zes = ze_raw.strip().lower() if isinstance(ze_raw, str) else ''
    if rds == 'n/a' or (rds == 'none' and zes == 'none'):
        abr = 'intern'
    elif rds == 'nicht nötig':
        abr = 'ohne'
    else:
        abr = 'rechnung'
    pid = f"p{year}_{c:03d}"
    col2id[c] = pid
    projects.append(dict(
        id=pid, jahr=year, sort=c, bereich=bereich, kategorie=kategorie, kunde=kunde, name=name,
        mwst=mwst, betrag=betrag, kosten=kosten, stundenOfferte=so, stundenZiel=num(wsv.cell(10, c).value),
        effort=int(num(wsv.cell(12, c).value)) if num(wsv.cell(12, c).value) else None,
        abrechnung=abr, rechnungGeplant=None, rechnungsdatum=rd, zahlungsdatum=ze,
        rechnungsadresse='', referenz='', rechnungsEmail='', rechnungstext='', notiz='\n'.join(notes), archiviert=False))

entries, n = [], 0
day = datetime.date(year, 1, 1)
leap = (year % 4 == 0 and year % 100 != 0) or year % 400 == 0
for r in range(first_data, last_data + 1):
    a = ws.cell(r, 1).value
    is_feb29 = isinstance(a, str) and a.strip().startswith('29.02')
    rowvals = {c: num(ws.cell(r, c).value) for c in range(2, last_col + 1)}
    if is_feb29 and not leap:
        assert not any(rowvals.values()), 'Stunden am 29.02. in einem Nicht-Schaltjahr'
        continue
    # Kontrolle: Monat/Tag der Excel-Zeile stimmt mit dem laufenden Datum ueberein
    if isinstance(a, datetime.datetime): assert (a.month, a.day) == (day.month, day.day), (r, a, day)
    for c, v in rowvals.items():
        if v:
            n += 1
            entries.append(dict(id=f"e{year}_{n:05d}", datum=day.isoformat(), projectId=col2id[c], stunden=v, notiz=''))
    day += datetime.timedelta(days=1)
assert day == datetime.date(year + 1, 1, 1), day

# Vorjahre aus 'Kumulierte Stunden' (Spalten G/H = kumuliert inkl. Zuschlag 1.1)
k = wv['Kumulierte Stunden']
history = {}
for col, y in ((7, year - 2), (8, year - 1)):
    cum = [num(k.cell(r, col).value) or 0.0 for r in range(2, 367)]
    daily = [round((cum[i] - (cum[i - 1] if i else 0.0)) / 1.1, 4) for i in range(len(cum))]
    history[str(y)] = daily

data = dict(version=1, quelle=f"Excel-Import {year}",
            settings=dict(mwstSatz=8.1, lohnFaktor=0.65, stundenZuschlag=1.1, pensum=80, stundenProTag=8.5, ferientage=25),
            projects=projects, entries=entries, history=history)
json.dump(data, open(out, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print(len(projects), 'Projekte,', len(entries), 'Eintraege,', round(sum(e['stunden'] for e in entries), 2), 'Stunden')
print({y: round(sum(v), 2) for y, v in history.items()})
